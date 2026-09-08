"""Render the shared slide model to PowerPoint."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from content import W, H, MONO
from slides import S

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
ANCH  = {"t": MSO_ANCHOR.TOP, "c": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}

def rgb(h):
    return RGBColor.from_string(h)

def no_line(shape):
    shape.line.fill.background()

def add_rect(sl, e):
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if e["radius"] else MSO_SHAPE.RECTANGLE
    s = sl.shapes.add_shape(shp_type, Inches(e["x"]), Inches(e["y"]),
                            Inches(e["w"]), Inches(e["h"]))
    if e["radius"]:
        # adjustment is a fraction of the shorter side
        try:
            s.adjustments[0] = min(0.5, e["radius"] / min(e["w"], e["h"]))
        except (IndexError, ZeroDivisionError):
            pass
    if e["fill"]:
        s.fill.solid(); s.fill.fore_color.rgb = rgb(e["fill"])
    else:
        s.fill.background()
    if e["line"]:
        s.line.color.rgb = rgb(e["line"]); s.line.width = Pt(e["lw"])
    else:
        no_line(s)
    s.shadow.inherit = False
    if e["text"] if "text" in e else False:
        pass
    return s

def add_text(sl, e):
    tb = sl.shapes.add_textbox(Inches(e["x"]), Inches(e["y"]),
                               Inches(e["w"]), Inches(e["h"]))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ANCH[e["va"]]
    p = tf.paragraphs[0]
    p.alignment = ALIGN[e["align"]]
    p.line_spacing = e["lh"]
    r = p.add_run(); r.text = e["text"]
    f = r.font
    f.size = Pt(e["size"]); f.bold = e["bold"]; f.italic = e["italic"]
    f.name = e["font"]; f.color.rgb = rgb(e["color"])
    return tb

def add_bullets(sl, e):
    tb = sl.shapes.add_textbox(Inches(e["x"]), Inches(e["y"]),
                               Inches(e["w"]), Inches(e["h"]))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(e["items"]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.24
        p.space_after = Pt(max(0, (e["gap"] - 0.22) * 72))
        r = p.add_run(); r.text = "•   " + item
        r.font.size = Pt(e["size"]); r.font.name = "Calibri"
        r.font.color.rgb = rgb(e["color"])
    return tb

def add_line(sl, e):
    from pptx.enum.shapes import MSO_CONNECTOR
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(e["x1"]),
                                Inches(e["y1"]), Inches(e["x2"]), Inches(e["y2"]))
    c.line.color.rgb = rgb(e["color"]); c.line.width = Pt(e["w"])
    return c

def add_circle(sl, e):
    s = sl.shapes.add_shape(MSO_SHAPE.OVAL, Inches(e["x"]), Inches(e["y"]),
                            Inches(e["d"]), Inches(e["d"]))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(e["fill"])
    no_line(s); s.shadow.inherit = False
    if e["text"]:
        tf = s.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = e["text"]
        r.font.size = Pt(e["size"]); r.font.bold = e["bold"]
        r.font.name = "Calibri"; r.font.color.rgb = rgb(e["color"])
    return s

def main(out):
    prs = Presentation()
    prs.slide_width = Inches(W); prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]
    for elems in S:
        sl = prs.slides.add_slide(blank)
        for e in elems:
            k = e["k"]
            if k == "rect":   add_rect(sl, e)
            elif k == "text": add_text(sl, e)
            elif k == "bullets": add_bullets(sl, e)
            elif k == "line": add_line(sl, e)
            elif k == "circle": add_circle(sl, e)
    prs.save(out)
    print("wrote", out, f"({len(S)} slides)")

if __name__ == "__main__":
    main(sys.argv[1])
